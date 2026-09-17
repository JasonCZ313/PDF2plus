"""PDF 转换模块：转图片 / 转 Word（支持 OCR）

OCR 方案：优先 RapidOCR + ONNX（模型随包内置，真正离线中英），
其次本机 Tesseract；均不可用时仅嵌入页面截图，不中断转换。
"""
import fitz
import os
import sys
import io
import tempfile
from typing import List, Optional, Callable, Tuple, Dict

# PyMuPDF pix.save() 原生支持的格式
_PYMUPDF_FORMATS = {"png", "jpg", "jpeg", "pnm", "pgm", "ppm", "pbm", "pam", "psd", "ps"}

# 转图片：MuPDF 超大 pixmap 会报 Overly large image 并静默返回全白图（本机实测约 2.78 亿像素起）
_MAX_EXPORT_PIXELS = 200_000_000
_MAX_EXPORT_EDGE = 16000
_MIN_EXPORT_DPI = 72
_MAX_EXPORT_DPI = 1200

# ---- OCR 引擎缓存（进程内只初始化一次） ----
_ocr_engine = None
_ocr_engine_type = None  # "rapidocr" | "tesseract" | None
_ocr_tried = False

# 内置 ONNX 模型文件名（与 assets/ocr_models、rapidocr 包内 models 对齐）
_OCR_MODEL_DET = "PP-OCRv6_det_small.onnx"
_OCR_MODEL_CLS = "ch_ppocr_mobile_v2.0_cls_mobile.onnx"
_OCR_MODEL_REC = "PP-OCRv6_rec_small.onnx"


def _save_pixmap(pix, output_path: str, fmt: str):
    """保存 pixmap：原生格式直接写；BMP/TIFF 等用 Pillow 兜底。"""
    if fmt.lower() in _PYMUPDF_FORMATS:
        pix.save(output_path)
    else:
        from PIL import Image
        img_data = pix.tobytes("png")
        img = Image.open(io.BytesIO(img_data))
        img.save(output_path, format=fmt.upper() if fmt.upper() != "JPG" else "JPEG")


def _safe_export_dpi(page_width_pt: float, page_height_pt: float, requested_dpi: int) -> int:
    """
    按页尺寸把 DPI 限制在 MuPDF 可稳定渲染的范围内，避免超限后得到全白图。
    """
    requested = max(_MIN_EXPORT_DPI, min(_MAX_EXPORT_DPI, int(requested_dpi)))
    width_pt = max(1.0, float(page_width_pt))
    height_pt = max(1.0, float(page_height_pt))
    # 像素边长 = pt * dpi / 72
    max_by_edge = int(_MAX_EXPORT_EDGE * 72.0 / max(width_pt, height_pt))
    # 总像素上限 → dpi <= 72 * sqrt(MAX / (w_pt * h_pt))
    max_by_pixels = int(((_MAX_EXPORT_PIXELS * 5184.0) / (width_pt * height_pt)) ** 0.5)
    return max(_MIN_EXPORT_DPI, min(requested, max_by_edge, max_by_pixels))


def _pixmap_looks_blank(pix, sample_stride: int = 97) -> bool:
    """
    抽样判断 pixmap 是否几乎全白。
    MuPDF 超限失败时典型表现：尺寸很大但像素全是 255。
    """
    if pix is None or pix.width < 1 or pix.height < 1:
        return True
    samples = pix.samples
    channel_count = int(pix.n)
    if channel_count < 1 or not samples:
        return True
    step = max(channel_count, channel_count * max(1, int(sample_stride)))
    nonwhite_hits = 0
    checked = 0
    for offset in range(0, len(samples), step):
        checked += 1
        if channel_count >= 3:
            if (
                samples[offset] < 250
                or samples[offset + 1] < 250
                or samples[offset + 2] < 250
            ):
                nonwhite_hits += 1
        else:
            if samples[offset] < 250:
                nonwhite_hits += 1
        if nonwhite_hits >= 3:
            return False
    return checked > 0 and nonwhite_hits == 0


def _page_has_exportable_content(page) -> bool:
    """页面是否应有可见内容（用于区分「真空白页」与「超限白图」）。"""
    try:
        if len((page.get_text("text") or "").strip()) >= 5:
            return True
    except Exception:
        pass
    try:
        if page.get_images():
            return True
    except Exception:
        pass
    try:
        if page.annots():
            return True
    except Exception:
        pass
    try:
        if page.get_drawings():
            return True
    except Exception:
        pass
    return False


def _render_page_for_export(page, dpi: int):
    """按 DPI 渲染单页；强制不透明底，并包含注释（笔迹/文本框）。"""
    matrix = fitz.Matrix(float(dpi) / 72.0, float(dpi) / 72.0)
    return page.get_pixmap(matrix=matrix, alpha=False, annots=True)


def _render_page_with_dpi_guard(page, requested_dpi: int) -> Tuple[object, int, bool]:
    """
    安全渲染：先按页限制 DPI；若仍全白且页上有内容，则逐步降 DPI 重试。

    :return: (pixmap, 实际 DPI, 是否相对请求值做过下调)
    """
    requested = max(_MIN_EXPORT_DPI, min(_MAX_EXPORT_DPI, int(requested_dpi)))
    safe_dpi = _safe_export_dpi(page.rect.width, page.rect.height, requested)
    was_reduced = safe_dpi < requested
    pixmap = _render_page_for_export(page, safe_dpi)

    if _page_has_exportable_content(page) and _pixmap_looks_blank(pixmap):
        retry_dpi = max(_MIN_EXPORT_DPI, min(safe_dpi - 100, int(safe_dpi * 0.75)))
        while retry_dpi >= _MIN_EXPORT_DPI:
            retry_pix = _render_page_for_export(page, retry_dpi)
            if not _pixmap_looks_blank(retry_pix):
                return retry_pix, retry_dpi, True
            if retry_dpi <= _MIN_EXPORT_DPI:
                break
            next_dpi = max(_MIN_EXPORT_DPI, retry_dpi - 100)
            if next_dpi == retry_dpi:
                break
            retry_dpi = next_dpi
        # 重试仍白：返回最后一次结果，由上层提示
        was_reduced = True
    return pixmap, safe_dpi, was_reduced


def _candidate_ocr_model_dirs() -> List[str]:
    """收集可能存放 OCR 模型的目录（开发态 / 打包态 / rapidocr 包内）。"""
    candidate_dirs: List[str] = []
    if getattr(sys, "frozen", False):
        meipass_dir = getattr(sys, "_MEIPASS", None)
        if meipass_dir:
            candidate_dirs.append(os.path.join(meipass_dir, "assets", "ocr_models"))
        # onedir：_internal/assets/ocr_models
        exe_dir = os.path.dirname(sys.executable)
        candidate_dirs.append(os.path.join(exe_dir, "assets", "ocr_models"))
        candidate_dirs.append(os.path.join(exe_dir, "_internal", "assets", "ocr_models"))
        candidate_dirs.append(
            os.path.join(os.path.dirname(exe_dir), "assets", "ocr_models")
        )
    else:
        project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        candidate_dirs.append(os.path.join(project_root, "assets", "ocr_models"))
    try:
        import rapidocr as rapidocr_pkg
        candidate_dirs.append(
            os.path.join(os.path.dirname(rapidocr_pkg.__file__), "models")
        )
    except Exception:
        pass
    return candidate_dirs


def _resolve_bundled_model_paths() -> Optional[Dict[str, str]]:
    """定位 det/cls/rec 三个 ONNX；全部存在才返回路径字典。"""
    required_files = {
        "det": _OCR_MODEL_DET,
        "cls": _OCR_MODEL_CLS,
        "rec": _OCR_MODEL_REC,
    }
    for model_dir in _candidate_ocr_model_dirs():
        resolved = {
            key: os.path.join(model_dir, filename)
            for key, filename in required_files.items()
        }
        if all(os.path.isfile(path) for path in resolved.values()):
            return resolved
    return None


def _init_ocr():
    """初始化 OCR：RapidOCR → Tesseract → None。"""
    global _ocr_engine, _ocr_engine_type, _ocr_tried
    if _ocr_engine is not None:
        return _ocr_engine, _ocr_engine_type
    if _ocr_tried:
        return None, None

    _ocr_tried = True

    # 1) RapidOCR（完整版默认，中英离线）
    try:
        from rapidocr import RapidOCR
        rapidocr_params = {"Global.log_level": "error"}
        model_paths = _resolve_bundled_model_paths()
        if model_paths:
            rapidocr_params["Det.model_path"] = model_paths["det"]
            rapidocr_params["Cls.model_path"] = model_paths["cls"]
            rapidocr_params["Rec.model_path"] = model_paths["rec"]
        _ocr_engine = RapidOCR(params=rapidocr_params)
        _ocr_engine_type = "rapidocr"
        return _ocr_engine, _ocr_engine_type
    except Exception:
        pass

    # 2) 本机 Tesseract（可选备用）
    try:
        import pytesseract
        try:
            pytesseract.get_tesseract_version()
        except Exception:
            for possible_path in [
                r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
            ]:
                if os.path.exists(possible_path):
                    pytesseract.pytesseract.tesseract_cmd = possible_path
                    break
            else:
                raise RuntimeError("Tesseract not found")
        _ocr_engine = pytesseract
        _ocr_engine_type = "tesseract"
        return _ocr_engine, _ocr_engine_type
    except Exception:
        pass

    return None, None


def _ocr_image(image_path: str) -> str:
    """对单张图片 OCR，返回多行文字；失败返回空串。"""
    engine, etype = _init_ocr()
    if engine is None:
        return ""
    try:
        if etype == "rapidocr":
            ocr_result = engine(image_path)
            text_lines = getattr(ocr_result, "txts", None) if ocr_result is not None else None
            if not text_lines:
                return ""
            return "\n".join(str(line) for line in text_lines if line)
        elif etype == "tesseract":
            import pytesseract
            return pytesseract.image_to_string(image_path, lang="chi_sim+eng")
    except Exception:
        return ""
    return ""


def pdf_to_images(
    pdf_path: str,
    output_dir: str,
    dpi: int = 300,
    fmt: str = "png",
    progress_callback: Optional[Callable[[int], None]] = None,
    cancel_check: Optional[Callable[[], bool]] = None,
    copy_original_pdf: bool = False,
    password: str = "",
    **kwargs,
) -> dict:
    """
    PDF 每页转换为高分辨率图片。

    超高 DPI 会按页自动限制，并在「有内容却渲成全白」时降 DPI 重试，
    避免 MuPDF Overly large image 静默输出白图。

    :return: dict — success / files / warning / dpi_requested / dpi_used_min
    """
    from app.pdf_security import open_pdf

    if kwargs.get("password"):
        password = kwargs.get("password") or password

    result = {
        "success": False,
        "files": [],
        "warning": None,
        "dpi_requested": int(dpi),
        "dpi_used_min": None,
    }
    output_files: List[str] = []
    reduced_notes: List[str] = []
    used_dpi_values: List[int] = []

    try:
        os.makedirs(output_dir, exist_ok=True)
        pdf = open_pdf(pdf_path, password or "")
        try:
            total = len(pdf)
            base_name = os.path.splitext(os.path.basename(pdf_path))[0]
            requested_dpi = max(_MIN_EXPORT_DPI, min(_MAX_EXPORT_DPI, int(dpi)))

            for page_index in range(total):
                if cancel_check and cancel_check():
                    break
                page = pdf[page_index]
                pixmap, used_dpi, was_reduced = _render_page_with_dpi_guard(
                    page, requested_dpi
                )
                used_dpi_values.append(used_dpi)
                if was_reduced:
                    reduced_notes.append(f"第{page_index + 1}页→{used_dpi}DPI")
                    # 降 DPI 后仍全白且页上有内容：记入警告，仍写出文件便于排查
                    if _page_has_exportable_content(page) and _pixmap_looks_blank(pixmap):
                        reduced_notes.append(f"第{page_index + 1}页可能仍异常偏白")

                from app.i18n import t
                output_name = t(
                    "out.image_page", base=base_name, page=page_index + 1, fmt=fmt
                )
                output_path = os.path.join(output_dir, output_name)
                _save_pixmap(pixmap, output_path, fmt)
                output_files.append(output_path)
                if progress_callback:
                    progress_callback(int((page_index + 1) / total * 100))
        finally:
            pdf.close()

        if copy_original_pdf:
            import shutil
            dest_pdf = os.path.join(output_dir, os.path.basename(pdf_path))
            if os.path.normpath(dest_pdf) != os.path.normpath(pdf_path):
                shutil.copy2(pdf_path, dest_pdf)

        result["files"] = output_files
        result["success"] = len(output_files) > 0
        if used_dpi_values:
            result["dpi_used_min"] = min(used_dpi_values)
        if reduced_notes:
            unique_notes = list(dict.fromkeys(reduced_notes))
            from app.i18n import t
            notes_text = "、".join(unique_notes[:8])
            if len(unique_notes) > 8:
                notes_text += "…"
            result["warning"] = t(
                "convert.dpi_auto_reduced",
                requested=requested_dpi,
                notes=notes_text,
            )
        return result
    except Exception as exc:
        result["files"] = output_files
        result["success"] = False
        from app.i18n import t
        result["warning"] = t("convert.images_error", error=str(exc))
        return result


def _page_is_scan_like(page) -> bool:
    """
    判断单页是否更像「扫描/纯图页」（需要 OCR + 嵌图）。

    规则（避免短标题页被误判）：
    1. 可提取文字 ≥ 30 → 文字页
    2. 无可主导版面的大图 → 文字页（含短标题、空白页）
    3. 最大图片面积 ≥ 页面积约 35%，且文字 < 30 → 扫描/图片页
    """
    text = page.get_text("text").strip()
    if len(text) >= 30:
        return False

    page_area = abs(page.rect)
    if page_area <= 0:
        return False

    max_image_area = 0.0
    try:
        # get_image_info 返回含 bbox 的字典列表，适合计算图片占版面比例
        for image_info in page.get_image_info():
            bbox = image_info.get("bbox")
            if not bbox:
                continue
            image_area = abs(fitz.Rect(bbox))
            if image_area > max_image_area:
                max_image_area = image_area
    except Exception:
        # 回退：仅能统计内嵌图数量时，有图且几乎无字则视为扫描页
        try:
            if page.get_images() and len(text) < 30:
                return True
        except Exception:
            return False
        return False

    if max_image_area <= 0:
        return False
    return (max_image_area / page_area) >= 0.35


def _classify_pages(pdf_path: str, cancel_check=None, password: str = "") -> Tuple[List[int], List[int], int]:
    """预扫描 PDF：将每页分类为 text_pages 或 image_pages。
    返回 (text_page_numbers, image_page_numbers, total_pages)。"""
    from app.pdf_security import open_pdf

    text_pages = []
    image_pages = []
    doc = open_pdf(pdf_path, password or "")
    total = len(doc)
    try:
        for page_index in range(total):
            if cancel_check and cancel_check():
                break
            page = doc[page_index]
            if _page_is_scan_like(page):
                image_pages.append(page_index + 1)
            else:
                text_pages.append(page_index + 1)
    finally:
        doc.close()
    return text_pages, image_pages, total


def _build_docx_from_pages(
    pdf_path: str,
    output_path: str,
    text_pages: List[int],
    image_pages: List[int],
    total: int,
    progress_callback=None,
    cancel_check=None,
    ocr_available: bool = False,
    password: str = "",
):
    """逐页构建 docx：文字页提取文本，图片页 OCR + 嵌入截图。"""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from app.pdf_security import open_pdf

    doc = open_pdf(pdf_path, password or "")
    document = Document()

    # 设置默认字体
    style = document.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(11)

    try:
        for i in range(total):
            if cancel_check and cancel_check():
                doc.close()
                return False

            page_num = i + 1
            page = doc[i]

            # 添加页码标题
            heading = document.add_heading(f'第 {page_num} 页', level=2)
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

            if page_num in text_pages:
                # ---- 文字页：直接提取文本 ----
                text = page.get_text("text").strip()
                if text:
                    for paragraph_text in text.split('\n'):
                        para_text = paragraph_text.strip()
                        if para_text:
                            p = document.add_paragraph(para_text)
                            p.style.font.size = Pt(11)
            else:
                # ---- 图片页：渲染 + OCR + 嵌入图片 ----
                matrix = fitz.Matrix(200 / 72, 200 / 72)  # 200 DPI，平衡清晰度与文件大小
                pix = page.get_pixmap(matrix=matrix)
                img_bytes = pix.tobytes("png")

                # 嵌入页面截图
                img_stream = io.BytesIO(img_bytes)
                try:
                    inline_shape = document.add_picture(img_stream, width=Inches(5.5))
                    last_paragraph = document.paragraphs[-1]
                    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    pass

                # OCR 识别
                if ocr_available:
                    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                        tmp.write(img_bytes)
                        tmp_path = tmp.name
                    try:
                        ocr_text = _ocr_image(tmp_path)
                        if ocr_text and ocr_text.strip():
                            note = document.add_paragraph("〔OCR 识别文字〕")
                            note.runs[0].font.size = Pt(9)
                            note.runs[0].font.color.rgb = RGBColor(100, 100, 100)
                            for line in ocr_text.strip().split('\n'):
                                line = line.strip()
                                if line:
                                    p = document.add_paragraph(line)
                                    p.style.font.size = Pt(10)
                    finally:
                        try:
                            os.unlink(tmp_path)
                        except Exception:
                            pass

            # 页间分隔
            if i < total - 1:
                document.add_page_break()

            if progress_callback:
                progress_callback(15 + int((i + 1) / total * 70))  # 15→85%

        doc.close()
        # 写入前确保父目录存在，避免 API/异常路径下 FileNotFound
        output_parent = os.path.dirname(output_path)
        if output_parent:
            os.makedirs(output_parent, exist_ok=True)
        document.save(output_path)
        return True

    except Exception:
        doc.close()
        raise


def pdf_to_word(
    pdf_path: str,
    output_path: str,
    progress_callback: Optional[Callable[[int], None]] = None,
    copy_original_pdf: bool = False,
    cancel_check: Optional[Callable[[], bool]] = None,
    password: str = "",
    **kwargs
) -> dict:
    """
    PDF 转 Word（对外入口）。

    默认在独立子进程中执行，避免 pdf2docx / OCR 原生崩溃带走主界面；
    子进程内通过环境变量 PDF2PLUS_WORD_WORKER=1 走真实转换实现。
    """
    if kwargs.get("password"):
        password = kwargs.get("password") or password
    # 已在隔离子进程中：直接转换
    if os.environ.get("PDF2PLUS_WORD_WORKER") == "1":
        return _pdf_to_word_impl(
            pdf_path,
            output_path,
            progress_callback=progress_callback,
            copy_original_pdf=copy_original_pdf,
            cancel_check=cancel_check,
            password=password or "",
        )
    # 主进程 / GUI 线程：拉起子进程，崩溃不影响主窗口
    return _pdf_to_word_via_subprocess(
        pdf_path,
        output_path,
        progress_callback=progress_callback,
        copy_original_pdf=copy_original_pdf,
        cancel_check=cancel_check,
        password=password or "",
    )


def _pdf_to_word_impl(
    pdf_path: str,
    output_path: str,
    progress_callback: Optional[Callable[[int], None]] = None,
    copy_original_pdf: bool = False,
    cancel_check: Optional[Callable[[], bool]] = None,
    password: str = "",
) -> dict:
    """
    实际转换逻辑（仅应在 WORD_WORKER 子进程或显式测试中调用）。

    - 纯文字：优先 pdf2docx；失败则回退纯文本抽取
    - 扫描/混合：有图片页才初始化 OCR
    """
    from app.pdf_security import materialize_unlocked_copy, probe_encryption

    result = {
        "success": False, "warning": None,
        "image_pages": [], "text_pages": [],
        "total_pages": 0, "ocr_used": False
    }

    try:
        from pdf2docx import Converter
    except ImportError:
        from app.i18n import t
        result["warning"] = t("convert.missing_pdf2docx")
        return result

    # pdf2docx 不能直接打开加密文件：需要时先解锁到临时副本
    working_path = pdf_path
    temp_unlocked = None
    try:
        _can_open, needs_password = probe_encryption(pdf_path)
        if needs_password or password:
            unlocked = materialize_unlocked_copy(pdf_path, password or "")
            if unlocked != pdf_path:
                working_path = unlocked
                temp_unlocked = unlocked
    except ValueError as exc:
        result["warning"] = str(exc)
        return result

    try:
        # ===== 阶段1: 预扫描 —— 智能分类页面 (0~10%) =====
        if progress_callback:
            progress_callback(1)

        text_pages, image_pages, total = _classify_pages(
            working_path, cancel_check, password=""
        )
        result["total_pages"] = total
        result["text_pages"] = text_pages
        result["image_pages"] = image_pages

        if cancel_check and cancel_check():
            return result

        if progress_callback:
            progress_callback(15)

        # 有扫描/图片页才加载 OCR，纯文字路径不碰 onnxruntime
        ocr_available = False
        if image_pages:
            _, ocr_type = _init_ocr()
            ocr_available = ocr_type is not None

        # 三种策略写盘前统一创建输出目录（含用户自建多级路径）
        output_parent = os.path.dirname(output_path)
        if output_parent:
            os.makedirs(output_parent, exist_ok=True)

        used_plain_fallback = False

        # ===== 阶段2: 根据分类选择最佳转换策略 =====
        if not image_pages:
            # ---- 纯文字 PDF → 优先 pdf2docx；失败回退纯文本抽取 ----
            pdf2docx_ok = False
            converter = None
            try:
                converter = Converter(working_path)
                converter.convert(output_path, start=0, end=None)
                pdf2docx_ok = os.path.isfile(output_path) and os.path.getsize(output_path) > 0
            except Exception:
                pdf2docx_ok = False
            finally:
                if converter is not None:
                    try:
                        converter.close()
                    except Exception:
                        pass
            if not pdf2docx_ok:
                # 删掉可能损坏的半成品，再走稳定的文本抽取路径
                try:
                    if os.path.isfile(output_path):
                        os.remove(output_path)
                except OSError:
                    pass
                success = _build_docx_from_pages(
                    working_path, output_path, text_pages, image_pages, total,
                    progress_callback, cancel_check, False, password=""
                )
                if not success:
                    return result
                used_plain_fallback = True
            if progress_callback:
                progress_callback(85)

        elif not text_pages:
            # ---- 纯图片 PDF → OCR + 嵌入截图 ----
            if ocr_available:
                result["ocr_used"] = True
            success = _build_docx_from_pages(
                working_path, output_path, text_pages, image_pages, total,
                progress_callback, cancel_check, ocr_available, password=""
            )
            if not success:
                return result

        else:
            # ---- 混合 PDF → 逐页智能处理 ----
            if ocr_available:
                result["ocr_used"] = True
            success = _build_docx_from_pages(
                working_path, output_path, text_pages, image_pages, total,
                progress_callback, cancel_check, ocr_available, password=""
            )
            if not success:
                return result

        # ===== 阶段3: 后处理 (85~100%) =====
        if copy_original_pdf:
            import shutil
            dest_pdf = os.path.join(
                os.path.dirname(output_path),
                os.path.basename(pdf_path)
            )
            if os.path.normpath(dest_pdf) != os.path.normpath(pdf_path):
                shutil.copy2(pdf_path, dest_pdf)
            if progress_callback:
                progress_callback(92)

        # ---- 构建警告/提示信息 ----
        from app.i18n import t
        if used_plain_fallback:
            result["warning"] = t("convert.word_fallback")
        elif image_pages:
            if ocr_available:
                ocr_info = t("convert.ocr_used")
            else:
                ocr_info = t("convert.ocr_missing")

            if len(image_pages) == total:
                result["warning"] = t(
                    "convert.all_image_pages", total=total, ocr=ocr_info
                )
            else:
                pages_str = "、".join(str(page_num) for page_num in image_pages[:5])
                if len(image_pages) > 5:
                    pages_str += t("convert.more_pages", count=len(image_pages))
                result["warning"] = t(
                    "convert.mixed_image_pages", pages=pages_str, ocr=ocr_info
                )

        result["success"] = True
        if progress_callback:
            progress_callback(100)

    except Exception as e:
        from app.i18n import t
        result["warning"] = t("convert.error", error=str(e))
        result["success"] = False
    finally:
        if temp_unlocked and os.path.isfile(temp_unlocked):
            try:
                os.remove(temp_unlocked)
            except OSError:
                pass

    return result


def _word_worker_command(job_json_path: str) -> List[str]:
    """构造转 Word 隔离子进程命令行（开发态跑 main.py，打包态跑当前 exe）。"""
    if getattr(sys, "frozen", False):
        return [sys.executable, "--task-pdf-to-word", job_json_path]
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main_script = os.path.join(project_root, "main.py")
    return [sys.executable, main_script, "--task-pdf-to-word", job_json_path]


def _pdf_to_word_via_subprocess(
    pdf_path: str,
    output_path: str,
    progress_callback: Optional[Callable[[int], None]] = None,
    copy_original_pdf: bool = False,
    cancel_check: Optional[Callable[[], bool]] = None,
    password: str = "",
) -> dict:
    """
    在独立子进程中执行转 Word。

    子进程若因原生库崩溃退出，主进程仍可读到失败结果并弹窗，避免整窗闪退。
    """
    import json
    import time
    import subprocess

    result = {
        "success": False,
        "warning": None,
        "image_pages": [],
        "text_pages": [],
        "total_pages": 0,
        "ocr_used": False,
    }
    from app.i18n import t
    result["warning"] = t("convert.word_worker_incomplete")
    work_dir = tempfile.mkdtemp(prefix="pdf2plus_word_")
    job_path = os.path.join(work_dir, "job.json")
    progress_path = os.path.join(work_dir, "progress.txt")
    result_path = os.path.join(work_dir, "result.json")
    cancel_path = os.path.join(work_dir, "cancel.flag")

    job_payload = {
        "pdf_path": pdf_path,
        "output_path": output_path,
        "password": password or "",
        "copy_original_pdf": bool(copy_original_pdf),
        "progress_path": progress_path,
        "result_path": result_path,
        "cancel_path": cancel_path,
    }
    try:
        with open(job_path, "w", encoding="utf-8") as job_file:
            json.dump(job_payload, job_file, ensure_ascii=False)

        command = _word_worker_command(job_path)
        # Windows 下隐藏控制台黑窗，避免用户以为又开了一个程序
        popen_kwargs = {
            "stdout": subprocess.DEVNULL,
            "stderr": subprocess.DEVNULL,
        }
        if os.name == "nt":
            startup_info = subprocess.STARTUPINFO()
            startup_info.dwFlags |= subprocess.STARTF_USESHOWWINDOW
            popen_kwargs["startupinfo"] = startup_info
            popen_kwargs["creationflags"] = getattr(subprocess, "CREATE_NO_WINDOW", 0)

        process = subprocess.Popen(command, **popen_kwargs)
        last_progress = -1
        while True:
            if cancel_check and cancel_check():
                try:
                    with open(cancel_path, "w", encoding="utf-8") as cancel_file:
                        cancel_file.write("1")
                except OSError:
                    pass
                # 给子进程一点时间自行收尾，再强制结束
                try:
                    process.wait(timeout=2.0)
                except Exception:
                    try:
                        process.kill()
                    except Exception:
                        pass
                result["warning"] = None
                result["success"] = False
                return result

            if progress_callback and os.path.isfile(progress_path):
                try:
                    with open(progress_path, "r", encoding="utf-8") as progress_file:
                        raw = progress_file.read().strip()
                    if raw:
                        value = int(float(raw))
                        if value != last_progress:
                            last_progress = value
                            progress_callback(max(0, min(100, value)))
                except Exception:
                    pass

            exit_code = process.poll()
            if exit_code is not None:
                break
            time.sleep(0.12)

        # 读取子进程写出的结果
        if os.path.isfile(result_path):
            try:
                with open(result_path, "r", encoding="utf-8") as result_file:
                    loaded = json.load(result_file)
                if isinstance(loaded, dict):
                    result.update(loaded)
            except Exception:
                pass

        # 子进程异常退出且没有成功标记 → 判定为原生崩溃或未捕获失败
        if not result.get("success"):
            if process.returncode not in (0, None):
                from app.i18n import t
                crash_hint = t("convert.word_worker_crashed")
                if not result.get("warning") or result.get("warning") == t(
                    "convert.word_worker_incomplete"
                ):
                    result["warning"] = crash_hint
                elif "子进程" not in str(result.get("warning")) and "worker" not in str(
                    result.get("warning")
                ).lower():
                    result["warning"] = f"{result['warning']}\n{crash_hint}"
        return result
    except Exception as exc:
        from app.i18n import t
        result["warning"] = t("convert.word_worker_start_fail", error=str(exc))
        result["success"] = False
        return result
    finally:
        # 尽力清理临时作业目录（输出 docx 不在此目录）
        try:
            import shutil
            shutil.rmtree(work_dir, ignore_errors=True)
        except Exception:
            pass


def run_pdf_to_word_worker(job_json_path: str) -> int:
    """
    无界面转 Word 子进程入口：读作业 JSON，写进度/结果文件。

    :return: 成功 0，失败 1（供主进程判断）
    """
    import json

    os.environ["PDF2PLUS_WORD_WORKER"] = "1"
    # 与主界面语言对齐，保证警告文案不是写死的中文
    try:
        from app.preferences import get as pref_get
        from app.i18n import set_language
        set_language(pref_get("ui_language", "zh"))
    except Exception:
        pass
    try:
        with open(job_json_path, "r", encoding="utf-8") as job_file:
            job = json.load(job_file)
    except Exception:
        return 1

    progress_path = str(job.get("progress_path") or "")
    result_path = str(job.get("result_path") or "")
    cancel_path = str(job.get("cancel_path") or "")

    def _write_progress(value: int) -> None:
        if not progress_path:
            return
        try:
            with open(progress_path, "w", encoding="utf-8") as progress_file:
                progress_file.write(str(int(value)))
        except OSError:
            pass

    def _cancel_check() -> bool:
        return bool(cancel_path and os.path.isfile(cancel_path))

    convert_result = _pdf_to_word_impl(
        str(job.get("pdf_path") or ""),
        str(job.get("output_path") or ""),
        progress_callback=_write_progress,
        copy_original_pdf=bool(job.get("copy_original_pdf")),
        cancel_check=_cancel_check,
        password=str(job.get("password") or ""),
    )
    try:
        if result_path:
            parent_dir = os.path.dirname(result_path)
            if parent_dir:
                os.makedirs(parent_dir, exist_ok=True)
            with open(result_path, "w", encoding="utf-8") as result_file:
                json.dump(convert_result, result_file, ensure_ascii=False)
    except Exception:
        return 1

    if _cancel_check():
        return 2
    return 0 if convert_result.get("success") else 1
